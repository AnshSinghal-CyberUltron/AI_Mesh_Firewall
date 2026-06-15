from docx import Document
from docx.shared import Pt

doc = Document()

title = doc.add_heading('QUALITY & SECURITY MANAGEMENT EVIDENCE', 0)
doc.add_paragraph('DRDO Technology Development Fund (TDF) – Expression of Interest (EOI) Submission')
doc.add_paragraph('Category: Risk & Compliance')
doc.add_paragraph('Reference: EOI.pdf Page 14, Section 14')
doc.add_paragraph('Owner: Yash')

sections = [
("Executive Summary",
"""This document demonstrates the organization's quality, security, AI governance, engineering assurance, and operational excellence capabilities. The objective is to provide evidence of mature processes, internationally aligned frameworks, and secure development practices that support the successful delivery of mission-critical projects for DRDO. The organization follows a structured approach covering Quality Management Systems (QMS), Information Security Management Systems (ISMS), AI Governance, Secure Software Development Lifecycle (Secure SDLC), DevSecOps, MLOps, Independent Verification & Validation (IV&V), and continuous improvement mechanisms."""),

("Quality Management System (QMS)",
"""The organization has established a Quality Management System that governs project execution, engineering processes, documentation standards, review mechanisms, and continual improvement activities.

Key quality processes include:
• Project Planning and Quality Objectives
• Requirements Management and Traceability
• Design Reviews and Architecture Assessments
• Configuration Management
• Change Management
• Internal Quality Audits
• Corrective and Preventive Actions (CAPA)
• Customer Feedback Management
• Continuous Improvement Programs

Quality reviews are conducted throughout the project lifecycle to ensure adherence to contractual, technical, security, and regulatory requirements."""),

("Information Security Management (ISO 27001 Alignment)",
"""The organization maintains a security framework aligned with ISO/IEC 27001 principles to protect information assets against cyber threats and operational risks.

Security domains covered include:
• Asset Management
• Access Control Management
• Identity and Privileged Access Management
• Cryptographic Controls
• Network Security
• Endpoint Security
• Security Monitoring
• Incident Response
• Vulnerability Management
• Business Continuity Management
• Disaster Recovery Planning
• Supplier Security Management

Security risks are identified, assessed, treated, monitored, and periodically reviewed through a formal risk management process."""),

("Artificial Intelligence Governance (ISO 42001 Alignment)",
"""The organization follows AI governance principles aligned with ISO/IEC 42001 to ensure trustworthy, explainable, transparent, and secure AI systems.

AI Governance controls include:
• AI Policy Framework
• AI Risk Assessment
• Dataset Quality Validation
• Model Explainability
• Bias and Fairness Testing
• Human Oversight Controls
• AI Security Reviews
• Ethical AI Governance
• AI Lifecycle Management
• Continuous Monitoring and Performance Evaluation

These controls ensure responsible development and deployment of AI-enabled systems."""),

("NIST AI Risk Management Framework",
"""The organization follows the NIST AI RMF methodology for identifying and managing AI-related risks.

Govern Function:
Establishes governance structures, accountability, and oversight mechanisms.

Map Function:
Identifies stakeholders, operational environments, assets, dependencies, and risk factors.

Measure Function:
Evaluates model performance, bias, security vulnerabilities, robustness, and reliability.

Manage Function:
Implements mitigation controls, continuous monitoring, remediation activities, and governance reporting.

This framework enables proactive management of AI risks throughout the project lifecycle."""),

("Secure Software Development Lifecycle (Secure SDLC)",
"""Security is integrated into every stage of software development.

Secure SDLC practices include:
• Security Requirements Definition
• Threat Modeling
• Secure Architecture Reviews
• Secure Coding Standards
• Code Reviews
• Static Application Security Testing (SAST)
• Dynamic Application Security Testing (DAST)
• Software Composition Analysis (SCA)
• Penetration Testing
• Vulnerability Remediation
• Security Acceptance Reviews

Security checkpoints are embedded throughout the development lifecycle to reduce vulnerabilities and strengthen product resilience."""),

("DevSecOps Framework",
"""The organization follows DevSecOps principles to integrate security into CI/CD pipelines.

DevSecOps capabilities include:
• Automated Security Scanning
• Source Code Security Analysis
• Dependency Scanning
• Container Security Assessments
• Infrastructure-as-Code Security Validation
• Secrets Management
• Compliance Monitoring
• Automated Deployment Controls
• Security Logging and Monitoring

These controls ensure continuous delivery while maintaining security assurance."""),

("MLOps Governance",
"""For AI and machine learning systems, MLOps controls are implemented to manage model lifecycle governance.

MLOps capabilities include:
• Dataset Versioning
• Model Version Control
• Automated Training Pipelines
• Validation Testing
• Performance Monitoring
• Drift Detection
• Secure Model Deployment
• Audit Logging
• Rollback Procedures
• Model Governance Controls

The framework ensures reliability, traceability, and operational stability of AI systems."""),

("Independent Verification & Validation (IV&V)",
"""Independent Verification and Validation activities provide assurance that systems meet project objectives and contractual requirements.

Verification activities include:
• Requirements Verification
• Design Verification
• Security Verification
• Performance Testing
• Reliability Testing

Validation activities include:
• User Acceptance Testing
• Operational Validation
• Compliance Validation
• Mission Readiness Assessments

Findings are documented and tracked through formal closure processes."""),

("Process Maturity and Continuous Improvement",
"""The organization promotes process maturity through structured governance, audits, reviews, and continuous improvement programs.

Improvement activities include:
• Internal Audits
• Root Cause Analysis
• Lessons Learned Programs
• Security Reviews
• Management Reviews
• KPI Monitoring
• Customer Satisfaction Programs

Continuous improvement initiatives are tracked through governance committees and management oversight."""),

("Evidence Matrix",
"""The following evidence can be provided in support of this submission:

• ISO 27001 Certificate
• Information Security Policies
• Risk Assessment Reports
• AI Governance Policies
• ISO 42001 Documentation
• Secure SDLC Procedures
• DevSecOps Framework Documents
• MLOps Governance Documentation
• Internal Audit Reports
• Security Testing Reports
• Penetration Testing Reports
• IV&V Reports
• Training Records
• Quality Management Procedures
• Business Continuity and Disaster Recovery Plans"""),

("Conclusion",
"""The organization has implemented a mature and comprehensive quality and security management ecosystem aligned with recognized international standards and industry best practices. These capabilities provide assurance that projects will be executed securely, efficiently, and in compliance with quality, security, and governance requirements expected for DRDO Technology Development Fund initiatives.""")
]

for heading, content in sections:
    doc.add_heading(heading, level=1)
    doc.add_paragraph(content)

path = "/mnt/data/DRDO_TDF_Detailed_Quality_Security_Management_Evidence.docx"
doc.save(path)
print(path)